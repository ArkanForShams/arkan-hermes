#!/usr/bin/env python3
"""Strip default outlines on all autoshapes (round 2, correct enum check)."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

path = "/home/shams/hermes-workspace/missions/mission-03-beverage/deck-meridian.pptx"
prs = Presentation(path)
n = 0
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            shape.line.fill.background()
            shape.shadow.inherit = False
            n += 1
prs.save(path)
print("stripped outlines on", n, "autoshapes")