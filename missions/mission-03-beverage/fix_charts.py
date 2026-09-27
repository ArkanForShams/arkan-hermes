#!/usr/bin/env python3
"""Debug + force chart series colors."""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

path = "/home/shams/hermes-workspace/missions/mission-03-beverage/deck-meridian.pptx"
prs = Presentation(path)
COLORS = ["0E4F52", "F5A623", "7FA982", "2E8C96"]
n = 0
for si, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_chart:
            chart = shape.chart
            print(f"slide {si}: chart_type={chart.chart_type}, plots={len(chart.plots)}")
            for pi, plot in enumerate(chart.plots):
                print(f"  plot {pi}: {len(plot.series)} series")
                for i, ser in enumerate(plot.series):
                    col = COLORS[i % len(COLORS)]
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
                    ser.format.line.color.rgb = RGBColor.from_string(col)
                    ser.format.line.width = Pt(2.75)
                    n += 1
prs.save(path)
print("styled", n, "series")